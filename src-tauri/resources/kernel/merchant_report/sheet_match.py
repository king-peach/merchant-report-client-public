# -*- coding: utf-8 -*-
"""多门店模式: 后台门店名 ↔ 模板 Sheet 的关键词匹配。

业务场景(用户口径):
  模板一店一 Sheet, Sheet 名是**地名关键词**(如 北辰 / 铁道学院 / 高桥);
  各后台的门店全名可能不一样, 但**会包含**这些地名关键词(如「长沙北辰三角洲店」);
  匹配不到的门店**不写入**即可(不报错、不自动建 Sheet —— 模板是权威)。

匹配优先级: 精确 > Sheet名(去后缀)是后台名子串(取最长=最具体) > 反向包含 > 配置关键词命中。
"""
import re

# 主数据表(汇总/流量/推广等), 不算门店表
DEFAULT_MAIN_SHEETS = ("Sheet1", "Sheet2", "Sheet3", "sheet1", "sheet2", "sheet3")
_STRIP = re.compile(r"[\s()（）\-—_·、,，.。/\\]+")
_SUFFIX = re.compile(r"(店|门店|分店|专营店|旗舰店|加盟店|有限公司|分公司)$")

# ---- 门店名 → Sheet 名的「地名关键字」抽取(用户口径) ----
# 长沙万家丽广场七楼店 → 万家丽广场七楼 ; 金盆岭理工大学店 → 金盆岭 ; 万家丽广场一楼店 → 万家丽广场一楼
DEFAULT_CITY_PREFIXES = ("湖南", "长沙", "湘潭", "株洲", "岳阳", "常德", "衡阳", "郴州", "益阳",
                         "娄底", "邵阳", "永州", "怀化", "张家界", "湘西", "宁乡", "浏阳", "望城",
                         "醴陵", "湘乡", "韶山", "汨罗", "临湘", "沅江", "资兴")
# 结尾的「场所/机构」噪声词 + 门店后缀: 命中就剥掉(循环剥, 直到剥不动)
# 注意: **不要**把「学院」当噪声词 —— 用户口径里「铁道学院」本身就是保留的地名关键字
# (实测教训: 剥掉"学院"会把「铁道学院店」抽成「铁道」, 与用户给的 Sheet 名不一致)。
DEFAULT_NOISE_WORDS = ("理工大学", "工业大学", "师范大学", "农业大学", "医科大学", "科技大学",
                       "财经大学", "大学", "中学", "小学", "幼儿园", "医院", "超市", "生活超市",
                       "便利店", "大市场", "农贸市场", "批发市场", "门店", "分店", "老店", "新店",
                       "专营店", "旗舰店", "加盟店", "店")
_QUALIFIER = {"老店", "新店", "分店", "旗舰店", "加盟店", "专营店", "总店", "门店"}
_PACK_RE = re.compile(r"[（(]([^（()）]{1,30})[)）]")
_BRAND_SEPS = ("·", "•", "▪", "∙")


def restyle_light(s):
    """轻量清洗(用于抽 Sheet 名): 只去空白与括号, **保留连字符/下划线/数字**(如 WJL-1楼)。"""
    return re.sub(r"[\s()（）\[\]【】]+", "", str(s or ""))


def extract_place(store_name, city_prefixes=None, noise_words=None, aliases=None):
    """从后台门店名里抽「地名关键字」当 Sheet 名(规则先这样定, 用户后续可细化)。

    处理链: 别名(优先) → **括号内定位词优先**(如「示例·…（金盆岭理工大学店）」→ 金盆岭)
            → 否则品牌分隔符后段/全名 → 去城市前缀(可多层) → 循环去结尾噪声词。
    aliases: {"金盆岭理工大学店": "金盆岭", ...} 直接指定, 优先级最高。
    """
    raw = str(store_name or "").strip()
    if not raw:
        return ""
    if aliases and raw in aliases:
        return str(aliases[raw])

    def _clean(s):
        s = restyle_light(s)
        for _ in range(4):                       # 去城市前缀(可能叠层, 如 湖南长沙)
            for c in (city_prefixes or DEFAULT_CITY_PREFIXES):
                c2 = restyle_light(c)
                if c2 and s.startswith(c2) and len(s) > len(c2) + 1:
                    s = s[len(c2):]
                    break
        changed = True
        while changed and len(s) > 2:            # 循环去结尾噪声词
            changed = False
            for w in (noise_words or DEFAULT_NOISE_WORDS):
                w2 = restyle_light(w)
                if w2 and s.endswith(w2) and len(s) - len(w2) >= 2:
                    s = s[: -len(w2)]
                    changed = True
                    break
        return restyle_light(s)

    # ① 括号内定位词优先(门店名常见「品牌（地名店）」结构)
    m = _PACK_RE.search(raw)
    if m:
        inner = _clean(m.group(1))
        if (inner and len(inner) >= 2
                and restyle_light(inner) not in {restyle_light(q) for q in _QUALIFIER}):
            if aliases and inner in aliases:
                return str(aliases[inner])
            return inner
    # ② 无括号/括号内是限定词 → 取品牌分隔符后段, 与全名比长度
    body = raw
    for sep in _BRAND_SEPS:
        if sep in body:
            body = body.split(sep)[-1]
            break
    body = _PACK_RE.sub("", body)
    cands = [c for c in (_clean(body), _clean(raw)) if c]
    best = max(cands, key=len, default=restyle_light(raw))
    if aliases and best in aliases:
        return str(aliases[best])
    return best


def norm(s):
    """归一化: 去空白/括号/标点, 便于子串比较。"""
    return _STRIP.sub("", str(s or ""))


def store_sheets(sheetnames, main=DEFAULT_MAIN_SHEETS):
    """挑出候选门店 Sheet: 排除主数据表。顺序保留模板原序。"""
    mainset = {norm(m) for m in main}
    return [sn for sn in sheetnames if norm(sn) and norm(sn) not in mainset]


def match_sheet(sheetnames, store_name, keywords=()):
    """把后台门店名匹配到某个 Sheet。返回 (sheet_name|None, 说明)。"""
    cand = [sn for sn in sheetnames if norm(sn)]
    name = norm(store_name)
    if not name:
        return None, "后台门店名为空"
    if not cand:
        return None, "模板里没有门店工作表"
    # 1) 精确
    for sn in cand:
        if norm(sn) == name:
            return sn, "精确匹配"
    # 2) Sheet 名(去「店」等后缀)是后台门店名的子串 → 取最长(最具体)
    hits = []
    for sn in cand:
        k = _SUFFIX.sub("", norm(sn))
        if len(k) >= 2 and k in name:
            hits.append((len(k), sn))
    if hits:
        hits.sort(key=lambda x: (-x[0], cand.index(x[1])))
        return hits[0][1], f"关键词包含(模板「{hits[0][1]}」⊂ 后台「{store_name}」)"
    # 3) 反向: 后台名(去后缀)是 Sheet 名的子串
    core = _SUFFIX.sub("", name)
    if len(core) >= 2:
        for sn in cand:
            if core in norm(sn):
                return sn, f"反向包含(后台「{store_name}」⊂ 模板「{sn}」)"
    # 4) 配置关键词(store.keyword / keywords): 同一关键词同时出现在 Sheet 名与门店名里
    for kw in keywords or ():
        k = norm(kw)
        if len(k) < 2:
            continue
        for sn in cand:
            if k in norm(sn) and k in name:
                return sn, f"配置关键词「{kw}」命中"
    return None, f"模板中无对应 Sheet(候选: {', '.join(cand[:8])})"


# ── 门店名映射表(模板里那张"区域/门店/淘宝闪购门店名称/京东门店名称/美团门店名称/门店id") ──
# 2026-09-29: 美团报「示例柠檬茶（长沙高桥店）」→ 关键词匹配能兜住, 但淘宝闪购那套名字
# 形状差得远(全/半角括号、·/? 之类的字符), 逐字匹配比猜稳。映射表是**用户自己维护的权威数据**,
# 只要模板里有就直接用; 没有才回退关键词匹配。
PLATFORM_COL_HINTS = {
    "taobao": ("淘宝", "闪购"),
    "meituan": ("美团",),
    "meituan_gj": ("美团",),      # 管家同样是美团系门店名
    "jd": ("京东",),
}


def _tidy(v):
    """门店名规范化: 去掉空白与常见装饰符号, 只留中日韩文字/字母/数字。"""
    t = re.sub(r"[\s]+", "", str(v or ""))
    return re.sub(r"[^0-9A-Za-z\u4e00-\u9fff]+", "", t)


def find_store_map_sheet(wb, platform, explicit=None):
    """找映射表位置 → (sheet, 表头行, 门店列, 平台列) 或 None。"""
    hints = PLATFORM_COL_HINTS.get(platform, ())
    order = ([explicit] if explicit and explicit in wb.sheetnames else []) + list(wb.sheetnames)
    for sn in order:
        ws = wb[sn]
        for hr in range(1, min(ws.max_row, 12) + 1):
            txt = []
            for c in range(1, min(ws.max_column, 24) + 1):
                v = ws.cell(row=hr, column=c).value
                txt.append(str(v).strip() if v is not None else "")
            if "门店" not in txt:
                continue
            i_sheet = txt.index("门店")
            i_name = next((i for i, t in enumerate(txt)
                           if "门店名称" in t and any(h in t for h in hints)), None)
            if i_name is not None:
                return sn, hr, i_sheet + 1, i_name + 1
    return None


def load_store_map(template_path, platform, explicit=None):
    """{规范化后端门店名: 模板 Sheet 名}。模板没有映射表 → 返回 {} (调用方回退关键词匹配)。"""
    try:
        import openpyxl
        wb = openpyxl.load_workbook(template_path, read_only=True, data_only=True)
    except Exception:
        return {}
    try:
        hit = find_store_map_sheet(wb, platform, explicit)
        if not hit:
            return {}
        sn, hr, c_sheet, c_name = hit
        ws = wb[sn]
        out = {}
        for r in range(hr + 1, min(ws.max_row, hr + 300) + 1):
            sheet = ws.cell(row=r, column=c_sheet).value
            nm = ws.cell(row=r, column=c_name).value
            if sheet is None or nm is None:
                continue
            key = _tidy(nm)
            if key:
                out[key] = str(sheet).strip()
        return out
    finally:
        try:
            wb.close()
        except Exception:
            pass


def lookup_store_map(store_map, store_name):
    """映射表里找: 规范化后精确 → 双向包含(后端名可能带后缀/前缀)。"""
    if not store_map:
        return None
    key = _tidy(store_name)
    if not key:
        return None
    if key in store_map:
        return store_map[key]
    for k, v in store_map.items():
        if k and (k in key or key in k):
            return v
    return None


def plan_matches(sheetnames, store_names, keywords=(), main=DEFAULT_MAIN_SHEETS,
                 store_map=None):
    """生成多门店写入计划。

    返回 dict(matched=[(后台门店, Sheet, 说明)], unmatched=[(后台门店, 原因)],
              idle=[Sheet...](模板里有、后台没有的门店), candidates=[Sheet...])
    """
    cand = store_sheets(sheetnames, main=main)
    kw = [k for k in (keywords or ()) if k]
    matched, unmatched, used = [], [], set()
    for name in store_names:
        sn = why = None
        # ① 优先用模板里的映射表(权威): 后端门店名 → 模板 Sheet 名
        _mapped = lookup_store_map(store_map, name)
        if _mapped and _mapped in sheetnames:
            sn, why = _mapped, "映射表(模板里的平台门店名)"
        elif _mapped:
            why = f"映射表指向「{_mapped}」但模板里没这张 Sheet"
        # ② 映射表没有/失效 → 回退关键词匹配
        if not sn:
            _sn2, _why2 = match_sheet(cand, name, kw)
            sn = _sn2
            # 映射表指向的表模板里没有时, 保留"映射表失效"这个更精确的解释(便于用户改模板)
            why = _why2 if _sn2 else (why or _why2)
        if sn:
            matched.append((name, sn, why))
            used.add(sn)
        else:
            unmatched.append((name, why))
    idle = [sn for sn in cand if sn not in used]
    return {"matched": matched, "unmatched": unmatched, "idle": idle, "candidates": cand}
