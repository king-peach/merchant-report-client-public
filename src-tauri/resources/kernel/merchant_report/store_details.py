# -*- coding: utf-8 -*-
"""《门店明细.xlsx》读取器 — 平台分组表头 → 编号/名称 → 模板 Sheet 名 映射。

用户定义格式(2026-10-09 实表):
  · 首个(或名为 data 的)工作表; 第 1 行 = 平台分组名(淘宝闪购/美团/美团管家/…, 可含注释文字;
    京东组允许留空 —— 用户注明"无法用编号导航+当前数据不需要单个"); 第 2 行 = 组内列头
    「门店名称」「门店唯一编号」; 另有「运营报表模板Sheet名」列(表头在 1 或 2 行, 值在数据行)。
  · 数据行从第 3 行起。
产出: {platform: {"by_id": {编号: sheet}, "by_name": {规范化名: sheet}}}
  platform ∈ taobao / meituan / meituan_gj / jd —— 与 PLATFORM_COL_HINTS 同一套键。
匹配纪律: 编号优先、名称兜底、对不上不猜(调用方留空+告警)。
"""
import re

from .sheet_match import _tidy

_GROUP_PLAT = (
    ("管家", "meituan_gj"),
    ("淘宝", "taobao"),
    ("闪购", "taobao"),
    ("京东", "jd"),
    ("美团", "meituan"),
)


def _platform_of(group):
    for kw, p in _GROUP_PLAT:
        if kw in group:
            return p
    return None


def load_store_details(path):
    """读门店明细表 → {platform: {"by_id": {...}, "by_name": {...}}}; 读不了 → {}。"""
    try:
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception:
        return {}
    try:
        sn = "data" if "data" in wb.sheetnames else wb.sheetnames[0]
        ws = wb[sn]
        rows = list(ws.iter_rows(min_row=1, max_row=min(ws.max_row, 500), values_only=True))
        if len(rows) < 2:
            return {}
        h1, h2 = rows[0], rows[1]
        ncol = max(len(h1), len(h2))
        get = lambda r, i: str(r[i]).strip() if i < len(r) and r[i] is not None else ""

        sheet_col = None
        for i in range(ncol):
            probe = get(h1, i) + get(h2, i)
            if "运营报表模板" in probe or ("Sheet" in probe and "名" in probe) or \
               re.search(r"模板.*[Ss]heet", probe):
                sheet_col = i
                break
        if sheet_col is None:
            return {}

        # 逐列: 组名(向前携带) + 列角色
        col_group, name_cols, id_cols = {}, [], []
        cur = ""
        for i in range(ncol):
            g = get(h1, i)
            if g and i != sheet_col:
                cur = g
            col_group[i] = cur
            role = get(h2, i)
            if "门店名称" in role:
                name_cols.append(i)
            elif "编号" in role or "ID" in role or "id" in role:
                id_cols.append(i)

        out = {}
        for i in name_cols:
            p = _platform_of(col_group[i])
            if not p:
                continue
            d = out.setdefault(p, {"by_id": {}, "by_name": {}})
            id_col = next((j for j in id_cols if col_group.get(j) == col_group[i]), None)
            for r in rows[2:]:
                nm = get(r, i)
                sheet = get(r, sheet_col)
                if not nm or not sheet:
                    continue
                d["by_name"].setdefault(_tidy(nm), sheet)
                if id_col is not None:
                    vid = get(r, id_col).split(".")[0]   # 防 Excel 转 float 丢精度样式 "33333333.0"
                    if vid:
                        d["by_id"].setdefault(vid, sheet)
        return out
    finally:
        try:
            wb.close()
        except Exception:
            pass
