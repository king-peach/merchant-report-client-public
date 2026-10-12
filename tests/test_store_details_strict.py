# -*- coding: utf-8 -*-
"""门店明细「严格模式」(需求 A2)单测 —— 用户口径 2026-10-12:

配了《门店明细》且能读时:
  ① 只写入能经明细映射到的门店, 其余不写入也不报错(日志"不在门店明细中 → 跳过", summary 计数);
  ② **不再自动新建 Sheet**(明细里的店模板里没有对应 Sheet → 跳过 + 明确日志);
  ③ 美团管家抓取循环只进明细里的门店(编号 12 位机构编码命中即保留; 名称兜底)。
未配置/读不出 → 一切保持旧行为(自动建表照旧)。本测试只用合成夹具。
"""
import os
import pathlib
import sys
import tempfile

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1] / "kernel"))
import openpyxl                                                                 # noqa: E402
from merchant_report import orchestrator as O                                   # noqa: E402
from merchant_report.platforms.meituan_gj import MeituanGjAdapter               # noqa: E402

KW = "岚霞路"


class FakeSink:
    def __init__(self):
        self.events = []

    def log(self, msg, level="info"):
        self.events.append({"type": "log", "msg": msg, "level": level})

    def need_human(self, *a, **kw):
        self.events.append({"type": "need_human"})


def logs(sink):
    return " ".join(e.get("msg", "") for e in sink.events if e["type"] == "log")


def mk_details(path):
    """门店明细: 美团组两家(岚霞路 编号 1001 / 绿地之窗 编号 1002)。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "data"
    ws.append(["美团", "", "运营报表模板Sheet名"])
    ws.append(["门店名称", "门店唯一编号", ""])
    ws.append(["岚霞路店", "1001", "岚霞路"])
    ws.append(["绿地之窗店", "1002", "绿地之窗"])
    wb.save(path)
    wb.close()
    return path


def mk_template(path, extra_sheets=("岚霞路",)):
    """主表 2 张(Sheet1/Sheet2) + 门店表; 注意 **没有**「绿地之窗」这张表。"""
    wb = openpyxl.Workbook()
    wb.active.title = "Sheet1"
    wb.create_sheet("Sheet2")
    for s in extra_sheets:
        wb.create_sheet(s)
    wb.save(path)
    wb.close()
    return path


def main():
    D = tempfile.mkdtemp(prefix="strict_")
    details = mk_details(f"{D}/门店明细.xlsx")
    names = ["岚霞路店", "绿地之窗店", "野店"]

    print("【1】未配置门店明细 → 旧行为: 匹配不到的门店照旧新建 Sheet(计划里有新表)")
    tpl_old = mk_template(f"{D}/old.xlsx")
    cfg = {"defaults": {}}
    store_old = {"template": tpl_old, "platform": "meituan", "keywords": []}
    sink = FakeSink()
    plan_old = O._plan_store_sheets(sink, store_old, names, cfg)
    assert O._create_missing_for(store_old, "meituan") is True, "未配置时 create_missing 必须为 True"
    made_old = [t[1] for t in plan_old["matched"]]
    assert KW in made_old, made_old
    assert any(t[2].startswith("新建") for t in plan_old["matched"] if t[0] == "野店"), plan_old["matched"]
    print(f"✅ 旧行为: matched={[(n, s) for n, s, _ in plan_old['matched']]}（野店按地名抽词建新表）")

    print("【2】配置门店明细 → 严格模式: 明细外的门店(野店)被过滤, 不写入也不报错")
    tpl = mk_template(f"{D}/t.xlsx")
    store = {"template": tpl, "platform": "meituan", "keywords": [], "store_details": details}
    sink = FakeSink()
    plan = O._plan_store_sheets(sink, store, names, cfg)
    matched_names = [t[0] for t in plan["matched"]]
    assert "野店" not in matched_names, plan["matched"]
    assert all("#" not in n for n in matched_names)
    assert "野店" in logs(sink) and "不在门店明细中 → 跳过" in logs(sink), logs(sink)
    assert "严格模式" in logs(sink), logs(sink)
    print(f"✅ 过滤明细外门店 ✓ matched={[(n, s) for n, s, _ in plan['matched']]}")

    print("【3】严格模式: 明细里指向的 Sheet 模板里没有 → 不建表 + 明确日志")
    assert O._create_missing_for(store, "meituan") is False, "严格模式必须禁止自动建表"
    names_after = openpyxl.load_workbook(tpl).sheetnames
    assert "绿地之窗" not in names_after, f"严格模式不该新建 Sheet: {names_after}"
    unmatched = [n for n, _ in plan["unmatched"]]
    assert "绿地之窗店" in unmatched, plan["unmatched"]
    assert "不新建 Sheet" in plans_log(plan, sink), "模板里没有对应 Sheet 要给可操作日志"
    # summary 体现跳过数量
    txt = O._strict_summary(plan, {"strict": True, "skipped": []})
    assert txt and txt.startswith("1 家不在门店明细中"), txt
    print(f"✅ 不建新表 ✓ unmatched={plan['unmatched'][:1]} · summary={txt}")

    print("【4】_details_scope: 编号优先(报表里带了门店编号时按编号命中)")
    sink = FakeSink()
    sc = O._details_scope(sink, store, "meituan", ["后台A", "后台B"],
                          ids_by_name={"后台A": "1001", "后台B": "9999"})
    assert sc["strict"] and sc["names"] == ["后台A"], sc
    print(f"✅ 编号命中 ✓ names={sc['names']} skipped={sc['skipped']}")

    print("【5】门店明细里没有该平台(如京东) → 不算严格模式(旧行为不变)")
    sink = FakeSink()
    sc2 = O._details_scope(sink, store, "jd", ["京东某店"])
    assert sc2["strict"] is False and sc2["names"] == ["京东某店"], sc2
    assert O._create_missing_for(store, "jd") is True, "该平台明细里没数据 → 保持旧行为"
    print("✅ 该平台无明细数据 → 不进入严格模式 ✓")

    print("【6】门店明细读不出 → 不进入严格模式(宁可不改行为, 也不误删门店)")
    sink = FakeSink()
    store_bad = {"template": tpl, "platform": "meituan", "store_details": f"{D}/不存在.xlsx"}
    assert O._create_missing_for(store_bad, "meituan") is True
    print("✅ 明细文件丢失 → 旧行为 ✓")

    print("【7】美团管家 allow 过滤: 编号(机构编码 12 位)优先, 名称兜底, 明细外不进店")
    from merchant_report.sheet_match import _tidy
    from merchant_report.platforms.meituan_gj import _gj_norm_name
    ad = MeituanGjAdapter(None, FakeSink(), None,
                          allow_name=["岚霞路店"], allow_id=["202000000002"])
    nm_norm = {_tidy(n) for n in ad.allow_name}      # 与 adapter 内部同一套归一化
    assert ad._allowed_store({"name": "岚霞路店", "sid": ""}, nm_norm) is True, "名称命中"
    assert ad._allowed_store({"name": "完全不认识的名字", "sid": "202000000002"}, nm_norm) is True, \
        "编号命中(卡片 sid 与明细同口径)"
    assert ad._allowed_store({"name": "野店", "sid": "999999999999"}, nm_norm) is False, "明细外必须跳过"
    # 未配置 allow → 不按门店明细过滤(与旧行为完全一致)
    ad2 = MeituanGjAdapter(None, FakeSink(), None)
    assert not ad2.allow_name and not ad2.allow_id
    print("✅ 管家 allow: 名称✓ / 编号✓ / 明细外✗ / 未配置=不过滤 ✓")

    # 名字规范化: 全角/空格差异也能命中(_gj_norm_name + _tidy)
    assert ad._allowed_store({"name": _gj_norm_name("岚 霞 路 店"), "sid": ""}, nm_norm) is True
    print("✅ 全角/空格差异仍能命中(归一化) ✓")

    print("\n✅ 门店明细严格模式(A2)全绿")


def plans_log(plan, sink):
    return logs(sink)


if __name__ == "__main__":
    main()
