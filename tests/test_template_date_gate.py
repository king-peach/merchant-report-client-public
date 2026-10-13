# -*- coding: utf-8 -*-
"""模板日期预检(需求 B)单测 —— 用户口径 2026-10-12:

跑流程时若发现**所有门店 Sheet 都没有我设置的那个日期** → 立刻停下问用户(继续/结束),
而不是先花几分钟采集完才发现。覆盖 4 种情形: 不存在→继续 / 不存在→结束 /
部分存在→不触发 / 无响应(无人值守)→继续+warn。

拆成两个可单测小函数: fill.template_date_scan(模板扫描) + orchestrator._template_date_gate(门)。
"""
import pathlib
import sys
import tempfile

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1] / "kernel"))
import openpyxl                                                          # noqa: E402
from merchant_report import fill                                         # noqa: E402
from merchant_report.orchestrator import _template_date_gate             # noqa: E402

MAINS = ("营业指标达成", "数据反馈看板")
DATE = "2026-10-11"


class FakeSink:
    def __init__(self):
        self.events = []

    def log(self, msg, level="info"):
        self.events.append({"type": "log", "msg": msg, "level": level})

    def need_human(self, action, message, **kw):
        self.events.append({"type": "need_human", "action": action, "message": message, **kw})

    def logs(self):
        return " ".join(e.get("msg", "") for e in self.events if e["type"] == "log")

    def panels(self):
        return [e for e in self.events if e["type"] == "need_human"]


class FakeHuman:
    def __init__(self, answers):
        self.answers, self.objs = list(answers), []

    def wait(self, obj, timeout_s=None):
        self.objs.append(obj)
        a = self.answers.pop(0) if self.answers else None
        return {"choice": a} if a else None


def mk(path, sheets):
    """sheets = [(sheet名, [A 列日期...]), ...]; 前 2 张固定为主表。"""
    wb = openpyxl.Workbook()
    wb.active.title = MAINS[0]
    wb.create_sheet(MAINS[1])
    for name, dates in sheets:
        ws = wb.create_sheet(name)
        ws["A5"] = "日期"
        for i, d in enumerate(dates, 6):
            ws.cell(i, 1).value = d               # 字符串日期(与真模板一致)
    wb.save(path)
    return path


def main():
    D = tempfile.mkdtemp(prefix="date_gate_")
    cfg = {"defaults": {}}

    print("【1】模板解析: 三吃(字符串/ datetime / Excel 序列号) + 排除主表/映射表")
    t = mk(f"{D}/t1.xlsx", [("岚霞路", [DATE]), ("绿地之窗", ["2026-10-10"])])
    assert fill.template_date_scan(t, DATE) == (2, 1), fill.template_date_scan(t, DATE)
    # Excel 序列号与 datetime 都要认(模板 A 列可能是日期格式, openpyxl 读成 datetime/数字)
    from datetime import datetime as _dt
    _serial = (_dt.strptime(DATE, "%Y-%m-%d") - fill.EPOCH).days
    wb = openpyxl.load_workbook(t)
    wb["绿地之窗"]["A6"] = _serial                # 序列号
    wb["岚霞路"]["A6"] = None
    wb.create_sheet("门店映射").append(["区域", "门店", "淘宝闪购门店名称"])   # _looks_like_store_map
    wb["门店映射"]["A5"] = DATE
    wb.save(t)
    assert fill.template_date_scan(t, DATE) == (2, 1), \
        f"映射表不算门店表: {fill.template_date_scan(t, DATE)}"
    print(f"✅ 扫描结果 {fill.template_date_scan(t, DATE)}（映射表被排除; 序列号 {_serial} 认得出）")

    print("【2】全新模板(没有可扫描门店 Sheet) → 跳过检查, 不打扰用户")
    t0 = mk(f"{D}/t0.xlsx", [])
    sink, human = FakeSink(), FakeHuman([])
    assert _template_date_gate(sink, human, t0, DATE, cfg) is True
    assert not human.objs and not sink.panels(), "全新模板不该弹确认"
    assert "没有可扫描的门店 Sheet" in sink.logs(), sink.logs()
    print("✅ 不弹面板 + 日志: " + sink.logs()[-60:])

    print("【3】所有门店表都没有该日期 + 用户选「继续跑流程」 → 继续")
    t2 = mk(f"{D}/t2.xlsx", [("岚霞路", ["2026-10-10"]), ("绿地之窗", ["2026-10-09"])])
    sink, human = FakeSink(), FakeHuman(["继续跑流程"])
    assert _template_date_gate(sink, human, t2, DATE, cfg) is True
    assert sink.panels() and sink.panels()[0]["choices"] == ["继续跑流程", "结束流程"]
    assert "已请求用户确认" in sink.logs(), sink.logs()
    print(f"✅ 弹面板 ✓ choices={sink.panels()[0]['choices']} · 选择继续 ✓")

    print("【4】所有门店表都没有该日期 + 用户选「结束流程」 → 中止")
    sink, human = FakeSink(), FakeHuman(["结束流程"])
    assert _template_date_gate(sink, human, t2, DATE, cfg) is False
    print("✅ 选择结束 → gate 返回 False(由 run_task 干净中止)")

    print("【5】部分门店表有该日期 → **不触发**(部分有部分没有属正常, 填表会自然补行)")
    t3 = mk(f"{D}/t3.xlsx", [("岚霞路", [DATE]), ("绿地之窗", ["2026-10-09"])])
    sink, human = FakeSink(), FakeHuman([])
    assert _template_date_gate(sink, human, t3, DATE, cfg) is True
    assert not human.objs and not sink.panels(), "部分存在不该弹确认"
    assert f"{DATE} 存在于 1/2 张门店表" in sink.logs(), sink.logs()
    print("✅ 不弹面板 + 日志: " + sink.logs()[-50:])

    print("【6】无人值守(human 返回 None) → 继续 + warn(保持自动化旧行为)")
    sink, human = FakeSink(), FakeHuman([])          # 会返回 None
    assert _template_date_gate(sink, human, t2, DATE, cfg) is True
    assert "无人值守/超时" in sink.logs(), sink.logs()
    print("✅ 无响应 → 继续 ✓ 日志: " + [l for l in sink.logs().split(" · ") if "无人" in l][-1][:80])

    print("【7】模板读不了(路径错) → 不打扰用户, 继续")
    sink, human = FakeSink(), FakeHuman([])
    assert _template_date_gate(sink, human, f"{D}/不存在.xlsx", DATE, cfg) is True
    assert not sink.panels()
    print("✅ 模板不可读 → 继续(检查失败不影响流程)")

    print("\n✅ 模板日期预检 + 确认门 全绿")


if __name__ == "__main__":
    main()
