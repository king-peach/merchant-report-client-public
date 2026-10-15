# -*- coding: utf-8 -*-
"""M5 守卫: 保存路径必须真的瘦身(防回归)。

背景: xlsx_slim.py 写得再对, 只要 save_wb 忘了调用(或重试阶梯里重复扫几百万格), 症状就会
原样回来 —— 输出还是 21.6MB。本测试**不看内部 API, 只看存盘产物**: 造一个"整行膨胀"的小
工作簿 → 走 fill.save_wb 真实保存路径 → 解包输出 xlsx, 数 sheet XML 里"越界且无值"的 <c>。

断言:
  ① 默认开 → 输出里越界空格式格 = 0(且"边界内"的格子还在, 不是把 sheet 清空了)
  ② 关掉(环境变量 / 配置) → 原样保存, 格子一个不少
  ③ 统计/日志: 有 log 出口时给「🧹 模板瘦身」一行, removed>0; 无 log 也不许炸
  ④ 内容格一字不差 + 行高/列宽活下来
"""
import os
import pathlib
import re
import sys
import tempfile
import zipfile

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1] / "kernel"))

import openpyxl                                                       # noqa: E402
from openpyxl.styles import Border, PatternFill, Side                # noqa: E402
from merchant_report import fill as fill                              # noqa: E402
from merchant_report import xlsx_slim as S                            # noqa: E402

FAILS = []
OLD_ENV = os.environ.pop("MRC_SLIM_XLSX", None)


def check(name, cond, extra=""):
    print(("  ✅ " if cond else "  ❌ ") + name + (f"  [{extra}]" if extra and not cond else ""))
    if not cond:
        FAILS.append(name)


def mk_inflated(path, title="门店A", lead=0):
    """内容 A1:J20, 之后铺满 K..(400) 列的**带样式空格**(WPS 整行套格式的等价物)。

    lead>0 时先垫 N 张空表 —— fill_template 的"主数据表 = 第 3 个 sheet"(main_sheet_in)
    要求模板至少 3 张表, 垫两张后目标表正好落在索引 2。
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "汇总" if lead else title
    if lead:
        for i in range(lead - 1):
            wb.create_sheet(f"垫表{i}")
        ws = wb.create_sheet(title)
    fill_ = PatternFill("solid", fgColor="FFF2CC")
    thin = Side(style="thin")
    bd = Border(left=thin, right=thin, top=thin, bottom=thin)
    for r in range(1, 21):
        for c in range(1, 11):
            ws.cell(row=r, column=c, value=r * 10 + c)
    for r in range(1, 61):
        for c in range(11, 401):
            cell = ws.cell(row=r, column=c)
            cell.fill, cell.border = fill_, bd
    ws.column_dimensions["C"].width = 24.0
    ws.row_dimensions[3].height = 33.0
    wb.save(path)
    wb.close()
    return path


_CELL_RE = re.compile(rb'<c r="([A-Z]{1,3})(\d+)"[^>]*?(/>|>.*?</c>)', re.S)


def blanks_outside(path, col_bound, row_bound):
    """解包 sheet XML, 数"越界且无值"的 <c>(无 <v>/<f>/<is> → 无内容)。"""
    n = 0
    with zipfile.ZipFile(path) as z:
        for name in z.namelist():
            if not re.match(r"xl/worksheets/sheet\d+\.xml$", name):
                continue
            xml = z.read(name)
            for m in _CELL_RE.finditer(xml):
                col_letters, row_s, body = m.group(1).decode(), int(m.group(2)), m.group(3)
                col = 0
                for ch in col_letters:
                    col = col * 26 + (ord(ch) - 64)
                if col <= col_bound and row_s <= row_bound:
                    continue
                if body in (b"/>", b"</c>") or not re.search(rb"<(v|f|is)[ >]", body):
                    n += 1
    return n


def test_default_on_slims():
    print("=== A. 默认开: 走 save_wb 真实保存路径 → 输出越界空格式格 = 0 ===")
    with tempfile.TemporaryDirectory() as d:
        tpl = mk_inflated(os.path.join(d, "t.xlsx"))
        wb = openpyxl.load_workbook(tpl)
        logs = []
        assert fill.save_wb(wb, tpl, log=logs.append) is True
        wb.close()
        st = S.slim_workbook(openpyxl.load_workbook(tpl))     # 只为拿边界值做断言
        check("输出越界无值格 = 0",
              blanks_outside(tpl, st["col_bound"], st["row_bound"]) == 0,
              str(blanks_outside(tpl, st["col_bound"], st["row_bound"])))
        check("内容区之外的带样式空格仍有保留(边界内未误剪)",
              blanks_outside(tpl, 10, 20) > 0, str(blanks_outside(tpl, 10, 20)))
        check("打了一行瘦身日志", any("🧹 模板瘦身" in m for m in logs), str(logs))
        check("日志含表数", any(re.search(r"（\d+ 张表）", m) for m in logs), str(logs))
        # 内容/格式活下来
        rb = openpyxl.load_workbook(tpl)
        rws = rb["门店A"]
        check("内容格在", rws["J20"].value == 210, repr(rws["J20"].value))
        check("行高保留", rws.row_dimensions[3].height == 33.0, str(rws.row_dimensions[3].height))
        check("列宽保留", rws.column_dimensions["C"].width == 24.0, str(rws.column_dimensions["C"].width))
        check("表名/表数不变", rb.sheetnames == ["门店A"], str(rb.sheetnames))
        rb.close()
        check("文件确实变小", os.path.getsize(tpl) < 400 * 1024, f"{os.path.getsize(tpl)}B")


def test_switch_off_keeps_everything():
    print("=== B. 开关 off(环境变量) → 原样保存, 一个格子都不少 ===")
    with tempfile.TemporaryDirectory() as d:
        tpl = mk_inflated(os.path.join(d, "t.xlsx"))
        before = zipfile.ZipFile(tpl).read("xl/worksheets/sheet1.xml")
        wb = openpyxl.load_workbook(tpl)
        os.environ["MRC_SLIM_XLSX"] = "0"
        try:
            logs = []
            assert fill.save_wb(wb, tpl, log=logs.append) is True
        finally:
            os.environ.pop("MRC_SLIM_XLSX", None)
        wb.close()
        after = zipfile.ZipFile(tpl).read("xl/worksheets/sheet1.xml")
        check("sheet XML 未变(未瘦身)", len(after) == len(before), f"{len(before)} → {len(after)}")
        check("没有瘦身日志", not any("🧹 模板瘦身" in m for m in logs), str(logs))
        wb2 = openpyxl.load_workbook(tpl)
        check("空格仍在(MRC_SLIM_XLSX=0 生效)", (60, 400) in wb2["门店A"]._cells)
        wb2.close()


def test_slim_false_kwarg():
    print("=== C. save_wb(slim=False) 显式关 / slim_enabled 传 cfg ===")
    with tempfile.TemporaryDirectory() as d:
        tpl = mk_inflated(os.path.join(d, "t.xlsx"))
        before = zipfile.ZipFile(tpl).read("xl/worksheets/sheet1.xml")
        wb = openpyxl.load_workbook(tpl)
        fill.save_wb(wb, tpl, slim=False)
        wb.close()
        check("slim=False → 未瘦身",
              len(zipfile.ZipFile(tpl).read("xl/worksheets/sheet1.xml")) == len(before))
        check("cfg defaults.slim_xlsx=false 判为关",
              S.slim_enabled(cfg={"defaults": {"slim_xlsx": False}}) is False)
        check("config._defaults() 默认 slim_xlsx=True(壳未配置也能生效)",
              __import__("merchant_report.config", fromlist=["x"])._defaults()
              ["defaults"].get("slim_xlsx") is True)


def test_idempotent_no_rescan_crash():
    print("=== D. 幂等: 连存两次不炸、第二次不再剪(重试阶梯里重复调也安全) ===")
    with tempfile.TemporaryDirectory() as d:
        tpl = mk_inflated(os.path.join(d, "t.xlsx"))
        wb = openpyxl.load_workbook(tpl)
        logs1, logs2 = [], []
        fill.save_wb(wb, tpl, log=logs1.append)
        fill.save_wb(wb, tpl, log=logs2.append)          # 同一个 wb 再存一次(已瘦过)
        wb.close()
        check("第一次有瘦身日志", any("🧹" in m for m in logs1), str(logs1))
        check("第二次 removed=0 → 不再打日志", not any("🧹" in m for m in logs2), str(logs2))
        check("重复保存后仍可载入", openpyxl.load_workbook(tpl).sheetnames == ["门店A"])


def test_fill_template_end_to_end():
    print("=== E. 端到端: fill_template 真实填表路径也瘦身 ===")
    with tempfile.TemporaryDirectory() as d:
        tpl = mk_inflated(os.path.join(d, "t.xlsx"), lead=2)   # 第 3 张表 = 主数据表
        logs = []
        res = fill.fill_template(tpl, "2026-10-15", {"C": 12345}, mode="top", log=logs.append)
        st = S.slim_workbook(openpyxl.load_workbook(tpl))
        check("写入成功(row 返回)", isinstance(res.get("row"), int), str(res))
        check("端到端输出越界无值格 = 0", blanks_outside(tpl, st["col_bound"], st["row_bound"]) == 0,
              str(blanks_outside(tpl, st["col_bound"], st["row_bound"])))
        check("端到端有瘦身日志", any("🧹 模板瘦身" in m for m in logs), str(logs))
        rb = openpyxl.load_workbook(tpl)
        rws = rb["门店A"]                                     # 主数据表 = 第 3 个 sheet
        check("写入值在盘上",
              any(rws[f"C{r}"].value == 12345 for r in range(5, 9)),
              str([rws[f"C{r}"].value for r in range(5, 9)]))
        check("内容/结构还在(J20 被 top 模式插行顶到 J21)",
              rws["J21"].value == 210, repr(rws["J21"].value))
        rb.close()


def main():
    print("M5 守卫: save_wb 保存路径必须瘦身")
    fill._sleep = lambda s: None                 # 本测试不该有真等待(占用重试用例不在此)
    test_default_on_slims()
    test_switch_off_keeps_everything()
    test_slim_false_kwarg()
    test_idempotent_no_rescan_crash()
    test_fill_template_end_to_end()
    print()
    if FAILS:
        print(f"❌ {len(FAILS)} 条失败: {FAILS}")
        return 1
    print("✅ 全部通过")
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    finally:
        if OLD_ENV is not None:
            os.environ["MRC_SLIM_XLSX"] = OLD_ENV
