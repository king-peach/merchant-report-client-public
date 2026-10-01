# -*- coding: utf-8 -*-
"""fill_many() 批量写入单测 —— 一店一 Sheet 的性能改造不能改变写入语义。

改造前: 每天每店各自调 fill_template = 一轮「读模板+写+存+回读」(京东 22 店 1 天实测 316s)。
改造后: fill.fill_many() 整任务一次读、一次写、一次存、一次回读。
本测试断言批量化后**行为与逐次调用等价**:
  ① 缺的门店 Sheet 会在批量里被建出来(复刻主表结构)
  ② 同一 Sheet 多天按日期升序写 → 最新日期在第一行(top 模式)
  ③ 公式列被还原(不是静态值)
  ④ 同日重跑覆盖同一行, 不新增行
  ⑤ 回读校验为空(值真的落盘了)
"""
import os
import sys
import tempfile

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "kernel"))

import openpyxl                                                     # noqa: E402
from merchant_report.fill import fill_many, fill_template, SHEET    # noqa: E402

FAILS = []


def check(name, cond, extra=""):
    print(("  ✅ " if cond else "  ❌ ") + name + (f"  [{extra}]" if extra and not cond else ""))
    if not cond:
        FAILS.append(name)


def make_template(path):
    """模仿真模板结构: 第 5 行指标名, 第 6 行起是数据区, B 列是公式列。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SHEET
    ws["A5"] = "日期"
    ws["B5"] = "星期"
    ws["C5"] = "营业额"
    ws["B6"] = '=TEXT(A6,"AAAA")'
    ws["E6"] = "=C6*2"
    wb.create_sheet("Sheet3")
    wb.save(path)
    return path


def main():
    d = tempfile.mkdtemp()
    tpl = make_template(os.path.join(d, "模板.xlsx"))

    print("【1】批量写入: 主表两天 + 门店 Sheet 两天(门店 Sheet 不存在 → 批量里新建)")
    jobs = [
        (None, "2026-09-20", {"C": 1, "D": 2}),
        ("金盆岭", "2026-09-20", {"C": 11}),
        ("金盆岭", "2026-09-21", {"C": 22}),
        (None, "2026-09-21", {"C": 3}),
    ]
    out = fill_many(tpl, jobs, mode="top")
    check("4 个写入任务全部返回", len(out) == 4, str(len(out)))
    check("空 write_set 的 job 被忽略", fill_many(tpl, [(None, "2026-09-19", None)]) == [],
          "空任务应返回 []")
    check("门店 Sheet 在批量里被建出来", "金盆岭" in {r[2]["sheet"] for r in out},
          str({r[2]["sheet"] for r in out}))
    check("每个结果都带 requested 字段", all("requested" in r[2] for r in out),
          str([r[2].get("requested") for r in out]))
    check("回读校验全部干净", all(not r[2]["reread_bad"] for r in out),
          str([r[2]["reread_bad"] for r in out]))

    print("【2】同日重跑覆盖同一行(不新增行)")
    wb = openpyxl.load_workbook(tpl)
    s2, sk = wb[SHEET], wb["金盆岭"]
    check("主表最新日期在第 6 行", str(s2["A6"].value)[:10] == "2026-09-21", str(s2["A6"].value))
    check("主表第二天在第 7 行", str(s2["A7"].value)[:10] == "2026-09-20", str(s2["A7"].value))
    check("门店 Sheet 最新日期在第 6 行", str(sk["A6"].value)[:10] == "2026-09-21", str(sk["A6"].value))
    check("门店 Sheet 第二天在第 7 行", str(sk["A7"].value)[:10] == "2026-09-20", str(sk["A7"].value))
    check("主表值落盘", s2["C6"].value == 3 and s2["C7"].value == 1, f"{s2['C6'].value}/{s2['C7'].value}")
    check("门店 Sheet 值落盘", sk["C6"].value == 22 and sk["C7"].value == 11,
          f"{sk['C6'].value}/{sk['C7'].value}")
    check("公式列被还原(不是静态值)", str(s2["B6"].value).startswith("=TEXT"),
          str(s2["B6"].value))
    check("门店 Sheet 公式也还原", str(sk["B6"].value).startswith("=TEXT"), str(sk["B6"].value))
    n_rows_before = s2.max_row
    wb.close()

    # 重跑必须显式指定主表(与编排层一致): 模板被跑过一轮后多出了门店 Sheet,
    # 此时"第 3 个 sheet"已是门店表 —— 自动解析在退化模板上不可靠, 由调用方钉名。
    fill_many(tpl, [(SHEET, "2026-09-21", {"C": 99})], mode="top")
    wb = openpyxl.load_workbook(tpl)
    check("同日重跑: 覆盖第 6 行而不是插新行",
          str(wb[SHEET]["A6"].value)[:10] == "2026-09-21" and wb[SHEET]["C6"].value == 99,
          f"{wb[SHEET]['A6'].value}/{wb[SHEET]['C6'].value}")
    check("行数没变", wb[SHEET].max_row == n_rows_before, f"{wb[SHEET].max_row} vs {n_rows_before}")
    check("没有重复建 Sheet", wb.sheetnames.count("金盆岭") == 1, str(wb.sheetnames))
    wb.close()

    print("【3】单次调用(旧路径)与批量结果一致 —— 同一份输入两种写法都落到同一格")
    tpl2 = make_template(os.path.join(d, "模板2.xlsx"))
    fill_template(tpl2, "2026-09-20", {"C": 5}, mode="top")
    fill_many(tpl2, [(None, "2026-09-20", {"C": 5})], mode="top")
    wb = openpyxl.load_workbook(tpl2)
    check("fill_template 与 fill_many 写同一行同一值",
          wb[SHEET]["C6"].value == 5 and str(wb[SHEET]["A6"].value)[:10] == "2026-09-20",
          f"{wb[SHEET]['C6'].value}")
    wb.close()

    print()
    if FAILS:
        print(f"❌ {len(FAILS)} 条失败: {FAILS}")
        return 1
    print("✅ 全部通过")
    return 0


if __name__ == "__main__":
    sys.exit(main())
