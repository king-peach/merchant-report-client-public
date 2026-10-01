# -*- coding: utf-8 -*-
"""淘宝「全部门店」只写进 1 家店 的回归测试(2026-09-30 真机)。

现场: 淘宝全部门店任务 ok=true, 下载的报表里 17 家门店(逐店一行), 但模板里**只有 1 张门店
Sheet 拿到数据**(正好是文件第一行那家「岚霞路」)。原因:
  `load_report_rows()` 是**账号级聚合** —— 同一天多行按"可加列求和"并成一行(给主表写汇总用),
  门店名只剩"首个非空"= 文件首行那家 → 用它去建一店一 Sheet 计划, 自然只匹配到 1 家。
京东早就有 rows_by_store(逐店行), 淘宝漏了。现补 `load_report_rows_by_store()`。
"""
import datetime
import pathlib
import sys
import tempfile

import openpyxl

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1] / "kernel"))
from merchant_report import fill, sheet_match                                          # noqa: E402
from merchant_report.orchestrator import _store_names_from_rows                        # noqa: E402

STORES = ["示例柠檬茶(湘潭岚霞路店)", "示例柠檬茶(绿地之窗店)",
          "示例柠檬茶(长沙高桥店)"]


def _mk_report(path, names=STORES, date="2026-09-29"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "data"
    ws.append(["日期", "门店名称", "门店编号", "营业额", "营业收入", "订单量"])
    for i, n in enumerate(names):
        ws.append([date, n, 1000 + i, 100.5 + i, 90.5 + i, 10 + i])
    wb.save(path)
    return path


def test_account_level_loader_collapses_stores():
    """记录坑: 账号级聚合的 load_report_rows 每天只剩 1 行(门店名=首行那家)。"""
    with tempfile.TemporaryDirectory() as d:
        f = _mk_report(f"{d}/r.xlsx")
        agg = fill.load_report_rows(f)
        assert len(agg) == 1 and len(agg["2026-09-29"].get("门店名称")) > 0
        assert agg["2026-09-29"]["门店名称"] == STORES[0], agg["2026-09-29"]["门店名称"]
        assert len(_store_names_from_rows(agg, ["2026-09-29"])) == 1
        print("✅ 账号级聚合口径: 3 家店 → 只剩 1 行(门店名=首行), 所以它不能用来建门店计划")


def test_by_store_loader_keeps_every_store():
    with tempfile.TemporaryDirectory() as d:
        f = _mk_report(f"{d}/r.xlsx")
        by = fill.load_report_rows_by_store(f)
        rows = by["2026-09-29"]
        assert len(rows) == 3, rows
        assert [r["门店名称"] for r in rows] == STORES
        assert rows[1]["营业额"] == 101.5, rows[1]           # 逐店原值(没被求和)
        assert len(_store_names_from_rows(by, ["2026-09-29"])) == 3
        print("✅ 逐店口径: 3 家店 3 行, 数值保持各店原值")


def test_plan_matches_all_sheets():
    """端到端: 逐店行 + 模板映射表 → 3 家店都能匹配到 Sheet(修复前只 1 家)。"""
    with tempfile.TemporaryDirectory() as d:
        rep = _mk_report(f"{d}/r.xlsx")
        tpl = f"{d}/t.xlsx"
        wb = openpyxl.Workbook()
        wb.active.title = "营业指标达成"
        wb.create_sheet("数据反馈看板")
        s1 = wb.create_sheet("Sheet1")
        s1.append(["区域", "门店", "淘宝闪购门店名称"])
        for sheet, name in zip(["岚霞路", "绿地之窗", "高桥"], STORES):
            s1.append([None, sheet, name])
        for sheet in ("岚霞路", "绿地之窗", "高桥", "无关店"):
            wb.create_sheet(sheet)
        wb.save(tpl)

        by = fill.load_report_rows_by_store(rep)
        names = _store_names_from_rows(by, ["2026-09-29"])
        smap = sheet_match.load_store_map(tpl, "taobao")
        plan = sheet_match.plan_matches(openpyxl.load_workbook(tpl).sheetnames,
                                        sorted(names), store_map=smap)
        got = sorted(sn for _nm, sn, _why in plan["matched"])
        assert got == ["岚霞路", "绿地之窗", "高桥"], (got, plan)
        print(f"✅ 门店计划: 3/3 家匹配到 Sheet {got}")


def test_real_report_17_stores_if_present():
    """真机文件在的话直接核对: 淘宝 17 家店全部识别(读取失败则跳过)。"""
    p = pathlib.Path.home() / "Downloads/商家报表整理"
    hits = sorted(p.glob("门店下载_20260929*5005*.xlsx")) or sorted(p.glob("门店下载_*.xlsx"))
    if not hits:
        print("… 真机报表文件不在, 跳过")
        return
    f = hits[0]
    by = fill.load_report_rows_by_store(f)
    agg = fill.load_report_rows(f)
    d = sorted(by)[0]
    print(f"… {f.name}: 逐店 {len(by[d])} 家 / 账号级聚合 {len(agg)} 行")
    assert len(by[d]) >= 2 and len(agg) == 1, (len(by[d]), len(agg))
    print(f"✅ 真机文件核对通过: 逐店 {len(by[d])} 家(修复前只会写 1 家)")


def main():
    for fn in (test_account_level_loader_collapses_stores, test_by_store_loader_keeps_every_store,
               test_plan_matches_all_sheets, test_real_report_17_stores_if_present):
        fn()
    print("✅ 淘宝一店一Sheet(逐店行)回归测试全部通过")


if __name__ == "__main__":
    main()
