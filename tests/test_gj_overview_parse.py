# -*- coding: utf-8 -*-
"""美团管家「营业概览」读数解析单测(纯文本, 不启动浏览器)。

背景(2026-09-30 真机, 第 1/31 家门店): 平台在营业概览里**新增了第 6 个渠道「自营外卖」**,
旧代码的渠道白名单只有 5 个 →
    渠道合计 2,040.81 vs 页面总量 2,074.81 → 差 34.00 = 1.64% > 1% 阈值
→ 交叉校验判「页面没刷完」→ **整店数据被弃用**(营业额只差 0.73%、订单量差 0.62% 所以没报,
只有营业收入被卡)。页面本身数据是完全自洽的(6 个渠道逐项相加 = 总量)。

盯住三条:
  ① 页面新增渠道 → 不能再丢掉整店数据(未识别渠道能解释差额就放行 + 告警);
  ② 未识别渠道要能写进模板 自营外卖 → HX/HY/HZ;
  ③ 真对不上(差额没人解释)仍然要拒绝, 不许把错数据写进表。

页面文本按 2026-09-30 诊断包截图逐字重建(诊断包以前不带页面文本, 只能靠截图抄)。
"""
import pathlib
import sys

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1] / "kernel"))
from merchant_report.platforms import gj_mapping as gm                          # noqa: E402

# 收入 167.30+77.50+1038.12+724.99+34.00+32.90 = 2074.81 ✓
# 营业 277.00+164.00+2287.40+1788.50+34.00+96.00 = 4646.90 ✓
# 单量 14+9+73+60+1+4 = 161 ✓
PAGE = """首页 经营分析 营业报表 菜品报表 收款报表 财务报表 运营报表
报表中心 报表中心首页 营业概览
营业日期: 2026/09/29 至 2026/09/29
查询时间范围: 2026/09/29 00:00:00-2026/09/29 23:59:59
查询 重置 图表模式 导出 取消常用 报表说明 重置
营业收入(元)
2,074.81
店内销售:167.30
自提销售:77.50
美团外卖:1,038.12
淘宝闪购:724.99
自营外卖:34.00
京东秒送:32.90
营业额(元)
4,646.90
店内销售:277.00
自提销售:164.00
美团外卖:2,287.40
淘宝闪购:1,788.50
自营外卖:34.00
京东秒送:96.00
优惠金额(元)
2,572.09
店内销售:109.70
自提销售:86.50
美团外卖:1,249.28
淘宝闪购:1,063.51
自营外卖:0.00
京东秒送:63.10
优惠占比
55.35%
店内销售:39.60%
自提销售:52.74%
美团外卖:54.62%
淘宝闪购:59.46%
自营外卖:0.00%
京东秒送:65.73%
订单量(单)
161
店内销售:14
自提销售:9
美团外卖:73
淘宝闪购:60
自营外卖:1
京东秒送:4
营业构成 订单分类构成 销售品项构成
品牌商-张泽来
01459
品牌商-张泽来
01459
"""

# 同一页面的另一种 innerText 顺序: 顶部指标卡片区在前, 明细构成区在后(旧代码就栽在这)
PAGE_CARDS_FIRST = """营业概览
营业日期: 2026/09/29 至 2026/09/29
营业收入(元) 2,074.81 营业额(元) 4,646.90 优惠金额(元) 2,572.09 优惠占比 55.35% 订单量(单) 161
""" + PAGE


def test_real_page_all_channels():
    parsed, meta = gm.scrape_overview_text(PAGE)
    assert not meta["crosscheck_failed"], meta["crosscheck_failed"]
    assert not meta["zero_data"]
    # 6 个渠道都解析出来了(含新渠道自营外卖)
    assert set(parsed) == {"店内销售", "自提销售", "美团外卖", "淘宝闪购", "自营外卖", "京东秒送"}, set(parsed)
    assert parsed["自营外卖"] == {"营业收入": 34.0, "营业额": 34.0, "订单量": 1.0}, parsed["自营外卖"]
    assert parsed["美团外卖"]["营业收入"] == 1038.12
    assert parsed["京东秒送"]["营业收入"] == 32.90
    print("✅ 真机页面文本: 6 渠道全部解析, 交叉校验通过")
    # 自营外卖 → 模板 HX/HY/HZ
    ws = gm.build_write_set(parsed)
    assert ws["HX"] == 34.0 and ws["HY"] == 34.0 and ws["HZ"] == 1.0, {k: ws[k] for k in ("HX", "HY", "HZ")}
    assert ws["HO"] == 277.0 and ws["HP"] == 167.30 and ws["HQ"] == 14.0
    print("✅ 自营外卖写入模板 HX/HY/HZ(34.00 / 34.00 / 1)")


def test_cards_row_first_still_works():
    """顶部卡片区排在明细区之前(旧实现取"第一次出现"→ 解析到没有渠道的卡片块)。"""
    parsed, meta = gm.scrape_overview_text(PAGE_CARDS_FIRST)
    assert not meta["crosscheck_failed"], meta["crosscheck_failed"]
    assert parsed["自营外卖"]["营业收入"] == 34.0
    assert parsed["店内销售"]["营业收入"] == 167.30
    probe = meta["probe"]["metrics"]["营业收入"]
    assert probe["candidates"] >= 2, probe       # 同一指标多处出现 → 逐处试
    print(f"✅ 指标多处出现(候选 {probe['candidates']} 个)仍取到自洽块: 营业收入 {probe['total']}")


def test_unknown_channel_reported_but_not_fatal():
    """页面再加新渠道(模板没列): 差额能被它解释 → 放行 + 记为未识别渠道(不再丢整店)。"""
    txt = PAGE.replace("京东秒送:32.90", "京东秒送:32.90\n抖音外卖:12.00").replace(
        "营业收入(元)\n2,074.81", "营业收入(元)\n2,086.81")
    parsed, meta = gm.scrape_overview_text(txt)
    assert not meta["crosscheck_failed"], meta["crosscheck_failed"]
    assert meta["unknown_channels"].get("抖音外卖") == 12.0, meta["unknown_channels"]
    assert "抖音外卖" not in parsed                       # 模板无列 → 不写
    print(f"✅ 未识别渠道被记账不致命: {meta['unknown_channels']}")


def test_scrambled_page_still_rejected():
    """真对不上(少了京东秒送那一行且没人解释差额) → 仍然必须拒绝, 防止写错数据。"""
    txt = PAGE.replace("京东秒送:32.90\n", "", 1)
    parsed, meta = gm.scrape_overview_text(txt)
    assert meta["crosscheck_failed"] == ["营业收入"], meta["crosscheck_failed"]
    detail = meta["probe"]["metrics"]["营业收入"]
    assert abs(detail["offset"] - 0.0157) < 0.002, detail
    print(f"✅ 差额无人解释 → 仍判失败(偏差 {detail['offset']:.4f}, 不写错数据)")


def test_noise_pairs_not_treated_as_channel():
    """页面其他「名称: 数值」(商户号等)不能被当成渠道, 否则差额会被噪声"解释"掉。"""
    txt = PAGE.replace("营业构成 订单分类构成 销售品项构成",
                       "营业构成 订单分类构成 销售品项构成 商户号:52759903 门店编号:01459")
    parsed, meta = gm.scrape_overview_text(txt)
    assert not meta["unknown_channels"], meta["unknown_channels"]
    assert not meta["crosscheck_failed"]
    print("✅ 商户号/门店编号等噪声对不参与渠道校验")


def test_zero_data_page_unchanged():
    """当天没数据(标签在、值全 0、构成块「暂无数据」) 仍走 zero_data, 不误报成结构变更。"""
    txt = ("营业概览 营业日期 至 查询 重置\n营业收入(元)\n0.00\n营业额(元)\n0.00\n"
           "优惠金额(元)\n0.00\n优惠占比\n0.00%\n订单量(单)\n0\n"
           "营业构成 订单分类构成 详情 暂无数据\n优惠构成 详情 暂无数据\n")
    parsed, meta = gm.scrape_overview_text(txt)
    assert parsed and meta["zero_data"] and meta["no_data_hint"], meta   # 全 0 = 真实值 → 产出结构照写
    assert all(v == 0 for d in parsed.values() for v in d.values()), parsed
    print("✅ 零数据页仍判 zero_data(不误报结构变更)")


def test_export_xlsx_keeps_extra_channel():
    """导出 xlsx 兜底路径: 自营外卖行同样要收(白名单换成 CHANNEL_SET)。"""
    import tempfile
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "营业"
    ws.append(["订单分类", "营业收入", "营业额", "优惠金额", "优惠占比", "订单量"])
    ws.append(["店内销售", 167.30, 277.00, 109.70, "39.60%", 14])
    ws.append(["自营外卖", 34.00, 34.00, 0.00, "0.00%", 1])
    ws.append(["京东秒送", 32.90, 96.00, 63.10, "65.73%", 4])
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        wb.save(f.name)
        parsed = gm.parse_biz_xlsx(f.name)
    assert parsed["自营外卖"]["营业收入"] == 34.00, parsed
    assert gm.build_write_set(parsed)["HY"] == 34.0
    print("✅ 导出 xlsx 兜底路径同样收到自营外卖")


def test_adapter_writes_store_instead_of_dropping_it():
    """端到端(假页面): 真机那段文本喂给 _scrape_one → 必须返回数据而不是 None。"""
    from merchant_report.platforms import meituan_gj as gj
    ad = gj.MeituanGjAdapter.__new__(gj.MeituanGjAdapter)
    logs = []
    ad.sink = type("S", (), {"log": lambda self, m, level="info": logs.append(m)})()
    ad._open_biz_overview = lambda page: page
    ad._set_gj_dates = lambda page, df, dt: None
    ad._wait_until = lambda cond, timeout_s=20, what="": True
    ad._text_of = lambda page: PAGE
    ad._last_err = ""
    data = ad._scrape_one(object(), "示例柠檬茶金盆岭理工大学店", "2026/09/29", "2026/09/29")
    assert data, f"整店数据被丢了: {ad._last_err}"
    assert data["自营外卖"] == {"营业收入": 34.0, "营业额": 34.0, "订单量": 1.0}, data["自营外卖"]
    assert not any("未映射渠道" in m for m in logs), logs     # 自营外卖已进模板 → 不该再告警
    print("✅ 端到端: 该店数据被采纳(不再因新增渠道整店弃用)")


def test_adapter_warns_on_truly_unmapped_channel():
    """模板真没有的渠道(如抖音外卖) → 采纳数据 + 明确告警, 用户能看见什么没写进去。"""
    from merchant_report.platforms import meituan_gj as gj
    ad = gj.MeituanGjAdapter.__new__(gj.MeituanGjAdapter)
    logs = []
    ad.sink = type("S", (), {"log": lambda self, m, level="info": logs.append(m)})()
    ad._open_biz_overview = lambda page: page
    ad._set_gj_dates = lambda page, df, dt: None
    ad._wait_until = lambda cond, timeout_s=20, what="": True
    ad._text_of = lambda page: PAGE.replace("京东秒送:32.90", "京东秒送:32.90\n抖音外卖:12.00").replace(
        "营业收入(元)\n2,074.81", "营业收入(元)\n2,086.81")
    ad._last_err = ""
    data = ad._scrape_one(object(), "某店", "2026/09/29", "2026/09/29")
    assert data, ad._last_err
    assert any("未映射渠道" in m and "抖音外卖" in m for m in logs), logs
    print("✅ 未映射渠道有明确告警: " + next(m for m in logs if "未映射渠道" in m))


def main():
    for fn in (test_real_page_all_channels, test_cards_row_first_still_works,
               test_unknown_channel_reported_but_not_fatal, test_scrambled_page_still_rejected,
               test_noise_pairs_not_treated_as_channel, test_zero_data_page_unchanged,
               test_export_xlsx_keeps_extra_channel, test_adapter_writes_store_instead_of_dropping_it,
               test_adapter_warns_on_truly_unmapped_channel):
        fn()
    print("✅ 管家营业概览解析(新增渠道不丢数据)单测全部通过")


if __name__ == "__main__":
    main()
