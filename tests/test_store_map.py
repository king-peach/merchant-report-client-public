# -*- coding: utf-8 -*-
"""门店名映射表单测(2026-09-29): 模板里那张「区域/门店/淘宝闪购门店名称/京东门店名称/美团门店名称/门店id」。

用户报「美团/淘宝都显示没找到门店, 不写入」→ 真因是关键词猜不准(全/半角括号、·/? 字符),
模板里的映射表本身就是权威数据 → 优先直接读它, 关键词匹配只作兜底。
"""
import pathlib
import shutil
import sys
import tempfile

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1] / "kernel"))
import openpyxl                                                             # noqa: E402
from merchant_report import sheet_match as SM                                # noqa: E402


def mk_template(path):
    wb = openpyxl.Workbook()
    wb.active.title = "营业指标达成"          # 第1个: 汇总(公式)
    wb.create_sheet("数据反馈看板")            # 第2个: 看板
    mp = wb.create_sheet("Sheet1")            # 第3个: 映射表
    mp["A2"], mp["B2"] = "区域", "门店"
    mp["C2"], mp["D2"], mp["E2"], mp["F2"] = "淘宝闪购门店名称", "京东门店名称", "美团门店名称", "门店id"
    mp["A3"], mp["B3"], mp["C3"] = "彭柳", "铁道", "示例柠檬茶(铁道学院店)"
    mp["D3"], mp["E3"], mp["F3"] = "示例柠檬茶(长沙铁道学院店)", "示例柠檬茶（铁道学院店）", 66666666
    mp["B4"], mp["C4"] = "高桥", "示例柠檬茶(长沙高桥店)"
    mp["D4"], mp["E4"], mp["F4"] = "示例柠檬茶(长沙高桥店)", "示例柠檬茶（长沙高桥店）", 28143799
    for nm in ("铁道", "高桥"):
        ws = wb.create_sheet(nm)
        ws["A5"] = "日期"
    wb.save(path)
    return path


def main():
    D = tempfile.mkdtemp(prefix="storemap_")
    tpl = mk_template(D + "/t.xlsx")

    for plat, want in (("meituan", "高桥"), ("taobao", "高桥"), ("jd", "高桥")):
        m = SM.load_store_map(tpl, plat)
        assert len(m) == 2, (plat, m)
        assert SM.lookup_store_map(m, {"meituan": "示例柠檬茶（长沙高桥店）",
                                       "taobao": "示例柠檬茶(长沙高桥店)",
                                       "jd": "示例柠檬茶(长沙高桥店)"}[plat]) == want, (plat, m)
        print(f"✅ [{plat}] 映射 {len(m)} 条, 该平台门店名 → 「{want}」")

    # 全角/半角、·/? 之类的差异不影响(规范化后匹配)
    m = SM.load_store_map(tpl, "meituan")
    # 规范化断言(全角/半角、空格、·?) —— **从映射表数据自身推导样例**, 不写死品牌名:
    # 写死会在公开脱敏版被一起替换 → 守卫失效 + 假失败(2026-10-01 踩过)。
    _keys = [str(k) for k in m]
    _base = next((k for k in _keys if "高桥" in k), None)
    if not _base:
        print("… 映射表里没有『高桥』门店名 → 跳过规范化断言")
    else:
        for variant in (_base, _base.replace("（", "(").replace("）", ")"),
                        _base.replace("?", "").replace("·", "").replace(" ", ""),
                        _base.replace("手打", " 手打 ")):
            assert SM.lookup_store_map(m, variant) == "高桥", variant
        print("✅ 全角/半角括号、空格、·? 等装饰符号差异都能对上")

    # plan_matches 走映射表, 说明里写明来源
    sheets = ["营业指标达成", "数据反馈看板", "Sheet1", "铁道", "高桥", "兰亭"]
    plan = SM.plan_matches(sheets, ["示例柠檬茶（长沙高桥店）"], ["高桥"], store_map=m)
    assert plan["matched"] == [("示例柠檬茶（长沙高桥店）", "高桥", "映射表(模板里的平台门店名)")], plan
    assert plan["unmatched"] == [] and "兰亭" in plan["idle"]
    print(f"✅ plan_matches 直接用映射表命中: {plan['matched'][0][1]}（{plan['matched'][0][2]}）")

    # 映射表缺失/失效 → 回退关键词匹配
    plan2 = SM.plan_matches(sheets, ["示例柠檬茶（长沙高桥店）"], ["高桥"], store_map={})
    assert plan2["matched"] and plan2["matched"][0][1] == "高桥", plan2
    assert "映射表" not in plan2["matched"][0][2]
    print(f"✅ 无映射表时回退关键词匹配: {plan2['matched'][0][2]}")
    plan3 = SM.plan_matches(sheets, ["某某店"], ["高桥"], store_map={"某某店": "不存在的表"})
    assert plan3["matched"] == [] and "不存在的表" in plan3["unmatched"][0][1], plan3
    print("✅ 映射表指向不存在的 Sheet → 不硬写, 落 unmatched 并说明")

    # 模板里没有映射表 → 返回 {} (调用方自然回退)
    p2 = D + "/plain.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "Sheet2"
    wb.create_sheet("高桥")
    wb.save(p2)
    assert SM.load_store_map(p2, "meituan") == {}
    print("✅ 模板没有映射表 → 返回 {}, 不影响老流程")

    shutil.rmtree(D, ignore_errors=True)
    print("✅ 门店名映射表测试全部通过")


if __name__ == "__main__":
    main()
